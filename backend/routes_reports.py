"""Real report exports: CSV, Excel (xlsx), PDF with SIES branding."""
import io
import csv
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from core import db, num
from auth import require
from engine import trial_balance, account_balances, INCOME, EXPENSE, ASSET, LIABILITY, EQUITY

router = APIRouter(prefix="/api")

INSTITUTION = "SIES College of Arts, Science & Commerce (Autonomous), Nerul, Navi Mumbai"


def fmt_inr(v):
    neg = v < 0
    s = f"{abs(v):,.2f}"
    # Indian grouping
    intpart, dec = s.split(".")
    intpart = intpart.replace(",", "")
    if len(intpart) > 3:
        last3 = intpart[-3:]
        rest = intpart[:-3]
        groups = []
        while len(rest) > 2:
            groups.insert(0, rest[-2:])
            rest = rest[:-2]
        if rest:
            groups.insert(0, rest)
        intpart = ",".join(groups) + "," + last3
    return ("-" if neg else "") + "\u20B9" + intpart + "." + dec


async def _report_data(report, fy):
    if report == "trial-balance":
        tb = await trial_balance(fy)
        cols = ["Code", "Account", "Type", "Debit", "Credit"]
        rows = [[r["code"], r["name"], r["type"], r["debit"], r["credit"]] for r in tb["rows"]]
        rows.append(["", "TOTAL", "", tb["total_debit"], tb["total_credit"]])
        return "Trial Balance", cols, rows, [3, 4]
    if report == "income-expenditure":
        bal, accounts = await account_balances(fy)
        cols = ["Code", "Account", "Category", "Amount"]
        rows = []
        for c, a in sorted(accounts.items()):
            v = bal.get(c, 0)
            if a["type"] == INCOME and v:
                rows.append([c, a["name"], "Income", v])
        for c, a in sorted(accounts.items()):
            v = bal.get(c, 0)
            if a["type"] == EXPENSE and v:
                rows.append([c, a["name"], "Expenditure", v])
        return "Income and Expenditure Statement", cols, rows, [3]
    if report == "fee-collection":
        cols = ["Roll No", "Student", "Program", "Assigned", "Paid", "Outstanding", "Status"]
        rows = []
        async for f in db.student_fees.find().sort("outstanding", -1).to_list(2000):
            rows.append([f["roll_no"], f["student_name"], f["program"],
                         num(f["assigned"]), num(f["paid"]), num(f["outstanding"]), f["status"]])
        return "Fee Collection Report", cols, rows, [3, 4, 5]
    if report == "budget-vs-actual":
        cols = ["Department", "Approved", "Actual", "Committed", "Available", "Utilization %"]
        rows = []
        async for b in db.budgets.find({"financial_year": fy}).sort("department", 1).to_list(100):
            rows.append([b["department_name"], num(b["approved"]), num(b["actual"]),
                         num(b["committed"]), num(b["available"]), b.get("utilization", 0)])
        return "Budget vs Actual Report", cols, rows, [1, 2, 3, 4]
    if report == "vendor-payables":
        cols = ["Invoice", "Vendor", "Department", "Amount", "Match", "Status"]
        rows = []
        async for i in db.invoices.find({"status": {"$ne": "paid"}}).to_list(2000):
            rows.append([i["invoice_no"], i["vendor_name"], i["department"],
                         num(i["amount"]), i["three_way_match"], i["status"]])
        return "Vendor Payables Report", cols, rows, [3]
    raise HTTPException(404, "Unknown report.")


REPORTS = ["trial-balance", "income-expenditure", "fee-collection",
           "budget-vs-actual", "vendor-payables"]


@router.get("/reports/{report}/export")
async def export_report(report: str, fmt: str = "csv", fy: str = "2025-26",
                        user: dict = Depends(require("reports", "export"))):
    if report not in REPORTS:
        raise HTTPException(404, "Unknown report.")
    title, cols, rows, money_cols = await _report_data(report, fy)

    if fmt == "csv":
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow([INSTITUTION])
        w.writerow([title, f"Financial Year {fy}"])
        w.writerow(cols)
        for r in rows:
            w.writerow(r)
        data = buf.getvalue().encode("utf-8")
        return _stream(data, "text/csv", f"{report}-{fy}.csv")

    if fmt == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        wb = Workbook()
        ws = wb.active
        ws.title = title[:28]
        ws.append([INSTITUTION])
        ws.append([f"{title}  |  Financial Year {fy}"])
        ws.append([])
        ws.append(cols)
        header_row = 4
        for cell in ws[header_row]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="0B2545")
        ws["A1"].font = Font(bold=True, size=12, color="EA580C")
        ws["A2"].font = Font(bold=True, size=11)
        for r in rows:
            ws.append([fmt_inr(c) if (isinstance(c, (int, float)) and i in money_cols) else c
                       for i, c in enumerate(r)])
        for col in ws.columns:
            width = max((len(str(c.value)) for c in col if c.value), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(width + 3, 45)
        out = io.BytesIO()
        wb.save(out)
        return _stream(out.getvalue(),
                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                       f"{report}-{fy}.xlsx")

    if fmt == "pdf":
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.units import mm
        from reportlab.platypus import (SimpleDocTemplate, Table, TableStyle,
                                        Paragraph, Spacer)
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        out = io.BytesIO()
        doc = SimpleDocTemplate(out, pagesize=landscape(A4),
                                topMargin=18 * mm, bottomMargin=16 * mm)
        styles = getSampleStyleSheet()
        h = ParagraphStyle("h", parent=styles["Title"], fontSize=15, textColor=colors.HexColor("#0B2545"))
        sub = ParagraphStyle("s", parent=styles["Normal"], fontSize=9, textColor=colors.HexColor("#EA580C"))
        meta = ParagraphStyle("m", parent=styles["Normal"], fontSize=8, textColor=colors.grey)
        elems = [Paragraph("SIES FinCore", sub), Paragraph(INSTITUTION, meta),
                 Spacer(1, 6), Paragraph(f"{title}", h),
                 Paragraph(f"Financial Year {fy}", meta), Spacer(1, 10)]
        table_data = [cols] + [
            [fmt_inr(c) if (isinstance(c, (int, float)) and i in money_cols) else str(c)
             for i, c in enumerate(r)] for r in rows]
        t = Table(table_data, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0B2545")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")]),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
            ("ALIGN", (max(money_cols) if money_cols else 1, 1), (-1, -1), "RIGHT"),
            ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        elems.append(t)
        doc.build(elems, onFirstPage=_pdf_footer, onLaterPages=_pdf_footer)
        return _stream(out.getvalue(), "application/pdf", f"{report}-{fy}.pdf")

    raise HTTPException(400, "Unsupported export format.")


def _pdf_footer(canvas, doc):
    from reportlab.lib.units import mm
    canvas.saveState()
    canvas.setFont("Helvetica", 7)
    canvas.setFillGray(0.5)
    canvas.drawString(18 * mm, 8 * mm, "Generated by SIES FinCore")
    canvas.drawRightString(doc.pagesize[0] - 18 * mm, 8 * mm, f"Page {doc.page}")
    canvas.restoreState()


def _stream(data, media, filename):
    return StreamingResponse(
        io.BytesIO(data), media_type=media,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'})
